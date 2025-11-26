import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Dict, List, Optional, Tuple

from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from compliance.enums import Frequency, InterestType, ParticularsType, PenaltyAmount
from compliance.models import ComplianceTaskMaster

SHEET_NAME = "Compliance Tasks Master"
ALLOWED_ROW_NUMBERS = frozenset(
    [
        8,
        12,
        13,
        14,
        15,
        16,
        19,
        20,
        21,
        22,
        25,
        26,
        27,
        29,
        30,
        31,
        40,
        41,
        42,
        43,
        45,
        47,
        48,
        52,
        53,
        54,
        55,
        56,
        57,
        58,
        59,
        61,
        63,
        64,
        69,
        70,
        71,
        72,
        73,
        74,
        75,
        76,
        77,
        78,
        79,
        80,
        81,
        82,
        83,
        84,
        85,
        86,
        88,
        90,
        92,
        93,
        94,
        95,
        96,
        97,
        98,
        99,
        100,
        101,
        102,
        103,
        110,
        111,
        120,
        121,
        122,
        124,
        125,
        129,
        130,
    ]
)

HEADER_FIELD_MAP = {
    "task id": "task_id",
    "act": "act",
    "particulars": "particulars",
    "due date": "due_date",
    "frequency": "frequency",
    "severity": "severity",
    "status": "status",
    "assignee": "assignee",
    "next due date": "next_due_date",
    "completed date": "completed_date",
    "penalty amount": "penalty_amount",
    "payment amount": "payment_amount",
    "payment reference": "payment_reference",
    "consequences": "consequences",
    "notes": "notes",
    "int": "interest_percentage",
    "penalty": "penalty",
    "late fees": "late_fee",
    "interest": "interest_amount",
}

DATE_FIELDS = {"due_date", "next_due_date", "completed_date"}
REQUIRED_FIELDS = {"act", "particulars", "frequency"}


class TaskImportError(Exception):
    """Raised when the uploaded sheet cannot be processed."""


@dataclass
class TaskImportResult:
    created: int = 0
    updated: int = 0
    skipped: List[Tuple[int, str]] = field(default_factory=list)


def import_compliance_tasks_from_excel(uploaded_file, acting_user) -> TaskImportResult:
    """Read the Excel file and create/update ComplianceTaskMaster records."""
    try:
        uploaded_file.seek(0)
    except (AttributeError, OSError):
        pass

    try:
        workbook = load_workbook(uploaded_file, data_only=True)
    except Exception as exc:  # pragma: no cover - openpyxl wraps errors
        raise TaskImportError("Unable to read the uploaded Excel file.") from exc

    if SHEET_NAME not in workbook.sheetnames:
        raise TaskImportError(
            f"'{SHEET_NAME}' sheet not found. Please upload the correct template."
        )

    sheet = workbook[SHEET_NAME]
    if sheet.max_row < 2:
        raise TaskImportError("The sheet does not contain any data rows.")

    header_map, header_row_index = _locate_header_row(sheet)
    if not header_map or header_row_index is None:
        raise TaskImportError(
            "Could not detect the header row. Place the column headers within the top 10 rows."
        )

    date1904 = bool(getattr(workbook.properties, "date1904", False))
    result = TaskImportResult()

    for row in sheet.iter_rows(min_row=header_row_index + 1):
        row_number = row[0].row
        if row_number not in ALLOWED_ROW_NUMBERS:
            continue

        row_payload = _extract_row_payload(row, header_map)
        if _row_is_empty(row_payload):
            continue

        try:
            task_payload = _transform_row_payload(row_payload, date1904)
        except ValueError as exc:
            result.skipped.append((row_number, str(exc)))
            continue

        missing_required = [
            field for field in REQUIRED_FIELDS if not task_payload.get(field)
        ]
        if missing_required:
            result.skipped.append(
                (row_number, f"Missing required values: {', '.join(missing_required)}")
            )
            continue

        try:
            created = _persist_task(task_payload, acting_user)
        except ValidationError as exc:
            result.skipped.append((row_number, _flatten_validation_error(exc)))
            continue
        except Exception as exc:  # pragma: no cover - defensive
            result.skipped.append((row_number, str(exc)))
            continue

        if created:
            result.created += 1
        else:
            result.updated += 1

    return result


def _build_header_map(header_row) -> Dict[int, str]:
    header_map: Dict[int, str] = {}
    for cell in header_row:
        header_key = _normalize_header(cell.value)
        field_name = HEADER_FIELD_MAP.get(header_key)
        if field_name:
            header_map[cell.col_idx] = field_name
    return header_map


def _normalize_header(value: Optional[str]) -> str:
    if value is None:
        return ""
    normalized = re.sub(r"[^0-9a-zA-Z]+", " ", str(value)).strip().lower()
    return normalized


def _extract_row_payload(row, header_map: Dict[int, str]) -> Dict[str, Optional[str]]:
    payload: Dict[str, Optional[str]] = {}
    for cell in row:
        field_name = header_map.get(cell.col_idx)
        if not field_name:
            continue
        payload[field_name] = cell.value
    return payload


def _row_is_empty(row_payload: Dict[str, Optional[str]]) -> bool:
    return all(value in (None, "", "NA", "N/A") for value in row_payload.values())


def _locate_header_row(sheet, max_scan_rows: int = 10):
    """
    Scan the first few rows to detect the header row.
    Returns a tuple of (header_map, header_row_index).
    """
    for row in sheet.iter_rows(min_row=1, max_row=max_scan_rows):
        header_map = _build_header_map(row)
        if not header_map:
            continue
        if REQUIRED_FIELDS.issubset(set(header_map.values())):
            return header_map, row[0].row
    return {}, None


def _normalize_choice_key(value: str) -> str:
    sanitized = str(value).strip()
    sanitized = sanitized.replace("–", "-").replace("—", "-")
    sanitized = re.sub(r"\s+", " ", sanitized)
    return sanitized.lower()


def _build_choice_lookup(choices):
    lookup = {}
    for stored_value, label in choices:
        normalized_value = _normalize_choice_key(stored_value)
        lookup[normalized_value] = stored_value
        normalized_label = _normalize_choice_key(label)
        lookup[normalized_label] = stored_value
    return lookup


FREQUENCY_LOOKUP = _build_choice_lookup(Frequency.choices)
PARTICULARS_LOOKUP = _build_choice_lookup(ParticularsType.choices)
PENALTY_LOOKUP = _build_choice_lookup(PenaltyAmount.choices)
PENALTY_ALIASES = {
    _normalize_choice_key(
        "Int - 18 % p.a. - Late Fees - ₹50 / day (₹20 NIL) - Penalty - ₹10,000 or 10 % of tax"
    ): PenaltyAmount.INT_18_PERCENT_PENALTY_10000,
    _normalize_choice_key(
        "Int - NA - Late Fees - ₹100/day (272A(2)(g)) - Penalty - ₹10k–₹1 lakh"
    ): PenaltyAmount.LATE_FEE_100_DAY_272A,
    _normalize_choice_key(
        "Int - 1% – 1.5% p.m. - Late Fees - NA - Penalty - ₹200/day (234E); ₹10k–₹1 lakh (271H)"
    ): PenaltyAmount.INT_1_TO_1_5_PERCENT_271H,
    _normalize_choice_key(
        "Int - NA - Late Fees - ₹100 per day - Penalty - No upper limit (per day fine)"
    ): PenaltyAmount.NO_UPPER_LIMIT_PENALTY,
}

FREQUENCY_ALIASES = {
    "every month (april - march)": Frequency.MONTHLY,
    "monthly / quarterly (tds)": Frequency.MONTHLY,
    "monthly (for statutory payments like tds, gst, advance tax, etc.)": Frequency.MONTHLY,
    "monthly (for pf/esi contributions)": Frequency.MONTHLY,
    "quarterly (end of each quarter)": Frequency.QUARTERLY,
    "quarterly (tds certificate)": Frequency.QUARTERLY,
    "quarterly (end of each quarter,quarterly (tds certificate)": Frequency.QUARTERLY,
    "event-based": Frequency.ANNUALLY,
    "continuous": Frequency.ANNUALLY,
    "annual": Frequency.ANNUALLY,
}


def _normalize_frequency(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    normalized = _normalize_choice_key(value)
    if normalized in FREQUENCY_ALIASES:
        return FREQUENCY_ALIASES[normalized]
    if "quarter" in normalized:
        return Frequency.QUARTERLY
    if any(token in normalized for token in ("month", "monthly")):
        return Frequency.MONTHLY
    if any(token in normalized for token in ("annual", "year")):
        return Frequency.ANNUALLY
    if "half" in normalized:
        return Frequency.HALF_YEARLY
    return FREQUENCY_LOOKUP.get(normalized, value)


INTEREST_LOOKUP = _build_choice_lookup(InterestType.choices)


def _normalize_interest_percentage(value: Optional[str]) -> Optional[str]:
    if value in (None, ""):
        return None
    normalized = _normalize_choice_key(value)
    if normalized in INTEREST_LOOKUP:
        return INTEREST_LOOKUP[normalized]
    sanitized = re.sub(r"[^\d\.\-]", "", value)
    if sanitized:
        try:
            numeric_value = float(sanitized)
            if numeric_value <= 1:
                numeric_value *= 100
            rounded = int(round(numeric_value))
            candidate = f"{rounded}%"
            candidate_key = _normalize_choice_key(candidate)
            if candidate_key in INTEREST_LOOKUP:
                return INTEREST_LOOKUP[candidate_key]
        except ValueError:
            pass
    return value


def _normalize_penalty_amount(value: Optional[str]) -> Optional[str]:
    if value in (None, ""):
        return None
    normalized = _normalize_choice_key(value)
    if normalized in PENALTY_LOOKUP:
        return PENALTY_LOOKUP[normalized]
    if normalized in PENALTY_ALIASES:
        return PENALTY_ALIASES[normalized]
    return value


def _normalize_particulars(value: Optional[str]) -> Optional[str]:
    if value in (None, ""):
        return None
    normalized = _normalize_choice_key(value)
    return PARTICULARS_LOOKUP.get(normalized, value)


def _transform_row_payload(row_payload, date1904) -> Dict[str, Optional[str]]:
    transformed: Dict[str, Optional[str]] = {}
    for field_name, raw_value in row_payload.items():
        if field_name in DATE_FIELDS:
            if raw_value in (None, "", "NA", "N/A"):
                transformed[field_name] = None
                continue
            try:
                transformed[field_name] = _coerce_excel_date(raw_value, date1904)
            except ValueError as exc:
                raise ValueError(f"{field_name}: {exc}") from exc
        else:
            coerced_value = _coerce_to_string(raw_value)
            if field_name == "frequency":
                coerced_value = _normalize_frequency(coerced_value)
            elif field_name == "interest_percentage":
                coerced_value = _normalize_interest_percentage(coerced_value)
            elif field_name == "penalty_amount":
                coerced_value = _normalize_penalty_amount(coerced_value)
            elif field_name == "particulars":
                coerced_value = _normalize_particulars(coerced_value)
            transformed[field_name] = coerced_value
    return transformed


def _coerce_excel_date(value, date1904) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        converted = from_excel(value, date1904=date1904)
        return converted.date()
    if isinstance(value, str):
        cleaned = value.strip()
        if not cleaned:
            return None
        candidates = {cleaned, cleaned.replace(".", "-"), cleaned.replace(".", "/")}
        formats = [
            "%d-%m-%Y",
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%d-%m-%y",
            "%d/%m/%y",
            "%d-%b-%Y",
            "%d-%b-%y",
            "%d %b %Y",
            "%d %B %Y",
        ]
        for candidate in candidates:
            for fmt in formats:
                try:
                    return datetime.strptime(candidate, fmt).date()
                except ValueError:
                    continue
    raise ValueError("Invalid date value.")


def _coerce_to_string(value) -> Optional[str]:
    if value in (None, "", "NA", "N/A"):
        return None
    if isinstance(value, str):
        cleaned = value.strip()
        return cleaned or None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _persist_task(task_payload: Dict[str, Optional[str]], acting_user) -> bool:
    task_id = task_payload.get("task_id")
    task = None
    if task_id:
        task = ComplianceTaskMaster.objects.filter(task_id=task_id).first()

    created = False
    if not task:
        task = ComplianceTaskMaster(task_id=task_id)
        created = True

    for field_name, value in task_payload.items():
        if field_name == "task_id" and not value:
            continue
        if value is not None:
            setattr(task, field_name, value)

    task.is_admin_created = True

    if created or not task.created_by:
        task.created_by = acting_user
    task.updated_by = acting_user
    task.save()
    return created


def _flatten_validation_error(error: ValidationError) -> str:
    if hasattr(error, "message_dict"):
        parts = []
        for field, messages in error.message_dict.items():
            joined = ", ".join(messages)
            parts.append(f"{field}: {joined}")
        return "; ".join(parts)
    return "; ".join(error.messages)
